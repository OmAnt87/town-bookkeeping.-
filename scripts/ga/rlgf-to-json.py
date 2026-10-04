#!/usr/bin/env python3
"""Reads the LOAD1 sheet of each Georgia RLGF report in data/raw/ga/rlgf/ into data/raw/ga/rlgf.json.

LOAD1 is the form's machine-readable copy: blocks named _R1.. (revenue), _E1.. (expenditures), _D1.. (debt) and _LOG1, each followed by a row of
field codes (31_1100, 1100A, SA_P1D, ...) and a row of values. Codes repeat across blocks (Part V
and Part X share function codes), so values are kept per block.

  pip install xlrd openpyxl && python3 scripts/ga/rlgf-to-json.py
"""
import json, os, re, sys

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
RAW = os.path.join(ROOT, 'data', 'raw', 'ga')


def rows_of(path):
    if path.endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if 'LOAD1' not in wb.sheetnames:
            return None
        return [list(r) for r in wb['LOAD1'].iter_rows(values_only=True)]
    import xlrd
    wb = xlrd.open_workbook(path)
    if 'LOAD1' not in wb.sheet_names():
        return None
    ws = wb.sheet_by_name('LOAD1')
    return [ws.row_values(i) for i in range(ws.nrows)]


def blocks_of(rows):
    out = {}
    for i, r in enumerate(rows):
        if r and isinstance(r[0], str) and re.fullmatch(r'_[A-Z]+\d+', r[0].strip()) and i + 2 < len(rows):
            head, vals = rows[i + 1], rows[i + 2]
            block = {}
            for h, v in zip(head, vals):
                if h in ('', None):
                    continue
                h = str(h).strip()
                if isinstance(v, str):
                    v = v.strip()
                    try:
                        v = float(v.replace(',', '')) if re.fullmatch(r'-?[\d,]+(\.\d+)?', v) else v
                    except ValueError:
                        pass
                block[h] = v if v != '' else None
            out[r[0].strip()[1:]] = block
    return out


def main():
    out, bad = {}, []
    files = sorted(os.listdir(os.path.join(RAW, 'rlgf')))
    for f in files:
        m = re.fullmatch(r'(\d+)_(\d{4})\.xlsx?', f)
        if not m:
            continue
        try:
            rows = rows_of(os.path.join(RAW, 'rlgf', f))
        except Exception as e:  # unreadable or not really a spreadsheet
            bad.append(f'{f}: {e}')
            continue
        if not rows:
            bad.append(f'{f}: no LOAD1 sheet')
            continue
        blocks = blocks_of(rows)
        if not blocks:
            bad.append(f'{f}: empty LOAD1')
            continue
        out.setdefault(m.group(1), {})[m.group(2)] = blocks
    with open(os.path.join(RAW, 'rlgf.json'), 'w') as fh:
        json.dump(out, fh)
    print(f'{sum(len(v) for v in out.values())} reports from {len(out)} governments; {len(bad)} unreadable')
    for b in bad[:30]:
        print('  ', b)


if __name__ == '__main__':
    sys.exit(main())
