#!/usr/bin/env python3
# Converts the Washington State Auditor's Financial Intelligence Tool (FIT) full extracts (one
# workbook per fiscal year, <year>.xlsx) into compact JSON for counties and cities/towns:
#   { accounts: { id: [logicalAccount, name, parentId] },
#     governments: { mcag: name },
#     years: { year: { mcag: { type, county, rows: [[accountId, fundTypeId, sectionId, amount]] } } } }
# Rows are summed by account, fund type and summary section.
#
#   python3 -I scripts/wa/fit-to-json.py <dir with <year>.xlsx> <out.json>
import json
import os
import re
import sys
from collections import defaultdict

import openpyxl

GOV_TYPES = {'06', '07'}  # county, city/town


def read_year(path, out):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data = {}
    for ws in wb.worksheets:
        title = ws.title
        if title.startswith('Schedule1'):
            it = ws.iter_rows(values_only=True)
            head = list(next(it))
            ix = {h: head.index(h) for h in ('MCAG', 'GovernmentTypeCode', 'County', 'Year', 'Amount', 'FundTypeId', 'BARSAccountId', 'FinancialSummarySection')}
            sums = defaultdict(float)
            meta = {}
            for r in it:
                if r[ix['GovernmentTypeCode']] not in GOV_TYPES:
                    continue
                mcag = r[ix['MCAG']]
                meta[mcag] = (r[ix['GovernmentTypeCode']], str(r[ix['County']]))
                sums[(mcag, r[ix['BARSAccountId']], r[ix['FundTypeId']], r[ix['FinancialSummarySection']])] += float(r[ix['Amount']] or 0)
            for (mcag, acct, fund, sec), amt in sums.items():
                g = data.setdefault(mcag, {'type': meta[mcag][0], 'county': meta[mcag][1], 'rows': []})
                g['rows'].append([acct, fund, sec, round(amt, 2)])
        elif title == 'Ref_Governments':
            it = ws.iter_rows(values_only=True)
            head = list(next(it))
            for r in it:
                out['governments'][r[head.index('MCAG')]] = r[head.index('CommonName')]
        elif title == 'Ref_AccountDesc':
            it = ws.iter_rows(values_only=True)
            head = list(next(it))
            for r in it:
                out['accounts'][r[head.index('Id')]] = [str(r[head.index('LogicalAccount')]), r[head.index('Name')], r[head.index('ParentId')]]
    return data


def main(src, dest):
    out = {'accounts': {}, 'governments': {}, 'years': {}}
    for f in sorted(os.listdir(src)):
        m = re.fullmatch(r'(\d{4})\.xlsx', f)
        if m:
            out['years'][m.group(1)] = read_year(os.path.join(src, f), out)
            print(f'{f}: {len(out["years"][m.group(1)])} governments', flush=True)
    with open(dest, 'w') as fh:
        json.dump(out, fh)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
