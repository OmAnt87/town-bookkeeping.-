#!/usr/bin/env python3
"""Reads the EDR revenue and expenditure workbooks in data/raw/fl/edr/ into data/raw/fl/edr.json.

Each workbook has one sheet per fiscal year. Account rows give the Uniform Accounting System
code, its name and amounts in twelve fund-type columns (general, special revenue, debt service,
capital projects, permanent, enterprise, internal service, custodial, pension, trust, private
purpose, component units). Category subtotal and total rows are skipped.

  pip install openpyxl && python3 scripts/fl/edr-to-json.py
"""
import json, os, re

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')
RAW = os.path.join(ROOT, 'data', 'raw', 'fl')
FIRST_YEAR = 2019
FUNDS = ['general', 'specialRevenue', 'debtService', 'capitalProjects', 'permanent', 'enterprise', 'internalService', 'custodial', 'pension', 'trust', 'privatePurpose', 'componentUnits']


def code_of(v):
    if isinstance(v, (int, float)):
        s = f'{v:.3f}'.rstrip('0').rstrip('.')
    else:
        s = str(v or '').strip()
    return s if re.fullmatch(r'\d{3}(\.\d{1,3})?', s) else None


def read(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    years = {}
    for name in wb.sheetnames:
        if not re.fullmatch(r'\d{4}', name.strip()) or int(name) < FIRST_YEAR:
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        fye = next((str(r[0]) for r in rows[:4] if r and r[0] and 'Fiscal Year' in str(r[0])), '')
        lines = []
        for r in rows:
            if not r:
                continue
            c = code_of(r[0])
            if not c:
                continue
            vals = [float(x) if isinstance(x, (int, float)) else float(str(x).replace(',', '') or 0) if x not in (None, '') and re.fullmatch(r'-?[\d,.]+', str(x)) else 0.0 for x in r[2:14]]
            lines.append([c, str(r[1] or '').strip(), dict(zip(FUNDS, vals))])
        years[name.strip()] = {'fye': fye, 'lines': lines}
    return years


def main():
    index = json.load(open(os.path.join(RAW, 'edr-index.json')))
    out, bad = {}, []
    for g in index:
        key = f"{g['kind']}-{g['slug']}"
        rec = {}
        for part in ('revenues', 'expenditures'):
            p = os.path.join(RAW, 'edr', f'{key}-{part}.xlsx')
            try:
                rec[part] = read(p)
            except Exception as e:
                bad.append(f'{key} {part}: {e}')
        out[key] = rec
    with open(os.path.join(RAW, 'edr.json'), 'w') as fh:
        json.dump(out, fh)
    print(f'{len(out)} governments; {len(bad)} unreadable workbooks')
    for b in bad[:20]:
        print('  ', b)


if __name__ == '__main__':
    main()
